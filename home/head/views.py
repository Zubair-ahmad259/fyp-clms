from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from student.models import Discipline
from .models import AdminProfile

@login_required
def admin_profile_list(request):
    """List all admin profiles with search and filter options"""
    admin_profiles = AdminProfile.objects.select_related('discipline').all()
    
    # Clear import errors from session if requested
    if request.GET.get('clear_errors'):
        if 'import_errors' in request.session:
            del request.session['import_errors']
    
    # Search functionality
    search_query = request.GET.get('search', '')
    if search_query:
        admin_profiles = admin_profiles.filter(
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(email__icontains=search_query) |
            Q(role__icontains=search_query) |
            Q(father_name__icontains=search_query) |
            Q(contact_number__icontains=search_query)
        )
    
    # Filter by role
    role_filter = request.GET.get('role', '')
    if role_filter:
        admin_profiles = admin_profiles.filter(role=role_filter)
    
    # Filter by discipline
    discipline_filter = request.GET.get('discipline', '')
    if discipline_filter:
        admin_profiles = admin_profiles.filter(discipline_id=discipline_filter)
    
    # Pagination
    paginator = Paginator(admin_profiles, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    # Get all disciplines for filter dropdown
    disciplines = Discipline.objects.all()
    
    context = {
        'page_obj': page_obj,
        'search_query': search_query,
        'role_filter': role_filter,
        'discipline_filter': discipline_filter,
        'disciplines': disciplines,
        'role_choices': AdminProfile._meta.get_field('role').choices,
        'import_errors': request.session.get('import_errors', []),
    }
    return render(request, 'admin_profile/admin_profile_list.html', context)

@login_required
def admin_profile_detail(request, pk):
    """View single admin profile details"""
    admin_profile = get_object_or_404(
        AdminProfile.objects.select_related('discipline'), 
        pk=pk
    )
    return render(request, 'admin_profile/admin_profile_detail.html', {
        'admin_profile': admin_profile
    })

@login_required
def admin_profile_create(request):
    """Create a new admin profile"""
    if request.method == 'POST':
        try:
            # Check if email already exists
            email = request.POST.get('email')
            if AdminProfile.objects.filter(email=email).exists():
                messages.error(request, f'Email "{email}" already exists!')
                return redirect('admin_profile:admin_profile_create')
            
            # Create admin profile
            discipline_id = request.POST.get('discipline')
            discipline = Discipline.objects.get(id=discipline_id) if discipline_id else None
            
            AdminProfile.objects.create(
                first_name=request.POST.get('first_name'),
                last_name=request.POST.get('last_name'),
                father_name=request.POST.get('father_name'),
                contact_number=request.POST.get('contact_number'),
                email=email,
                discipline=discipline,
                role=request.POST.get('role'),
                address=request.POST.get('address')
            )
            
            messages.success(request, 'Admin profile created successfully!')
            return redirect('admin_profile:admin_profile_list')
                
        except Exception as e:
            messages.error(request, f'Error creating admin profile: {str(e)}')
    
    # GET request - show form
    disciplines = Discipline.objects.all()
    context = {
        'disciplines': disciplines,
        'role_choices': AdminProfile._meta.get_field('role').choices,
    }
    return render(request, 'admin_profile/admin_profile_form.html', context)

@login_required
def admin_profile_update(request, pk):
    """Update an existing admin profile"""
    admin_profile = get_object_or_404(AdminProfile, pk=pk)
    
    if request.method == 'POST':
        try:
            # Check if email is being changed and if it's unique
            new_email = request.POST.get('email')
            if new_email != admin_profile.email and AdminProfile.objects.filter(email=new_email).exists():
                messages.error(request, f'Email "{new_email}" already exists!')
                return redirect('admin_profile:admin_profile_update', pk=pk)
            
            # Update admin profile
            discipline_id = request.POST.get('discipline')
            admin_profile.discipline = Discipline.objects.get(id=discipline_id) if discipline_id else None
            
            admin_profile.first_name = request.POST.get('first_name')
            admin_profile.last_name = request.POST.get('last_name')
            admin_profile.father_name = request.POST.get('father_name')
            admin_profile.contact_number = request.POST.get('contact_number')
            admin_profile.email = new_email
            admin_profile.role = request.POST.get('role')
            admin_profile.address = request.POST.get('address')
            admin_profile.save()
            
            messages.success(request, 'Admin profile updated successfully!')
            return redirect('admin_profile:admin_profile_detail', pk=admin_profile.pk)
                
        except Exception as e:
            messages.error(request, f'Error updating admin profile: {str(e)}')
    
    # GET request - show form with current data
    disciplines = Discipline.objects.all()
    context = {
        'admin_profile': admin_profile,
        'disciplines': disciplines,
        'role_choices': AdminProfile._meta.get_field('role').choices,
    }
    return render(request, 'admin_profile/admin_profile_form.html', context)

@login_required
def admin_profile_delete(request, pk):
    """Delete an admin profile"""
    admin_profile = get_object_or_404(AdminProfile, pk=pk)
    
    if request.method == 'POST':
        try:
            email = admin_profile.email
            admin_profile.delete()
            messages.success(request, f'Admin profile for "{email}" deleted successfully!')
            return redirect('admin_profile:admin_profile_list')
        except Exception as e:
            messages.error(request, f'Error deleting admin profile: {str(e)}')
    
    return render(request, 'admin_profile/admin_profile_confirm_delete.html', {
        'admin_profile': admin_profile
    })

@login_required
def import_admin_profiles(request):
    """Import admin profiles from Excel file"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        
        # Check file extension
        if not excel_file.name.endswith(('.xlsx', '.xls')):
            messages.error(request, 'Please upload an Excel file (.xlsx or .xls)')
            return redirect('admin_profile:admin_profile_list')
        
        try:
            # Load workbook
            wb = openpyxl.load_workbook(excel_file)
            sheet = wb.active
            
            # Get headers from first row
            headers = {}
            for col_idx, cell in enumerate(sheet[1], 1):
                if cell.value:
                    headers[str(cell.value).strip().lower()] = col_idx
            
            # Required columns
            required_columns = ['first_name', 'last_name', 'father_name', 'email', 'role']
            
            # Check if all required columns exist
            missing_columns = [col for col in required_columns if col not in headers]
            if missing_columns:
                messages.error(request, f'Missing required columns: {", ".join(missing_columns)}')
                return redirect('admin_profile:admin_profile_list')
            
            # Process data rows (starting from row 2)
            success_count = 0
            error_count = 0
            errors = []
            
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                # Skip empty rows
                if not any(row):
                    continue
                    
                try:
                    # Extract data using headers
                    first_name = row[headers['first_name']-1] if headers['first_name'] <= len(row) else None
                    last_name = row[headers['last_name']-1] if headers['last_name'] <= len(row) else None
                    father_name = row[headers['father_name']-1] if headers['father_name'] <= len(row) else None
                    email = row[headers['email']-1] if headers['email'] <= len(row) else None
                    role = row[headers['role']-1] if headers['role'] <= len(row) else None
                    
                    # Optional fields
                    contact_number = ''
                    if 'contact_number' in headers and headers['contact_number'] <= len(row):
                        contact_number = row[headers['contact_number']-1] or ''
                    
                    address = ''
                    if 'address' in headers and headers['address'] <= len(row):
                        address = row[headers['address']-1] or ''
                    
                    discipline_name = ''
                    if 'discipline' in headers and headers['discipline'] <= len(row):
                        discipline_name = row[headers['discipline']-1] or ''
                    
                    # Validate required fields
                    if not all([first_name, last_name, father_name, email, role]):
                        errors.append(f"Row {row_idx}: Missing required fields")
                        error_count += 1
                        continue
                    
                    # Clean data
                    first_name = str(first_name).strip()
                    last_name = str(last_name).strip()
                    father_name = str(father_name).strip()
                    email = str(email).strip().lower()
                    role = str(role).strip()
                    
                    # Check if email already exists
                    if AdminProfile.objects.filter(email=email).exists():
                        errors.append(f"Row {row_idx}: Email '{email}' already exists")
                        error_count += 1
                        continue
                    
                    # Validate role
                    valid_roles = [role[0] for role in AdminProfile._meta.get_field('role').choices]
                    if role not in valid_roles:
                        errors.append(f"Row {row_idx}: Invalid role '{role}'. Must be one of: {', '.join(valid_roles)}")
                        error_count += 1
                        continue
                    
                    # Get discipline
                    discipline = None
                    if discipline_name:
                        discipline_name = str(discipline_name).strip()
                        discipline = Discipline.objects.filter(name__iexact=discipline_name).first()
                        if not discipline:
                            errors.append(f"Row {row_idx}: Discipline '{discipline_name}' not found")
                            error_count += 1
                            continue
                    
                    # Create admin profile
                    AdminProfile.objects.create(
                        first_name=first_name,
                        last_name=last_name,
                        father_name=father_name,
                        contact_number=str(contact_number).strip() if contact_number else '',
                        email=email,
                        discipline=discipline,
                        role=role,
                        address=str(address).strip() if address else ''
                    )
                    
                    success_count += 1
                        
                except Exception as e:
                    errors.append(f"Row {row_idx}: {str(e)}")
                    error_count += 1
            
            # Show summary message
            if success_count > 0:
                messages.success(request, f'Successfully imported {success_count} admin profiles!')
            if error_count > 0:
                messages.warning(request, f'Failed to import {error_count} profiles. Check errors below.')
                request.session['import_errors'] = errors[:50]  # Store first 50 errors in session
            
        except Exception as e:
            messages.error(request, f'Error processing Excel file: {str(e)}')
    
    return redirect('admin_profile:admin_profile_list')

@login_required
def download_sample_excel(request):
    """Download sample Excel template for admin profile import"""
    # Create a new workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Admin Profiles"
    
    # Define headers
    headers = [
        'first_name', 'last_name', 'father_name', 'email', 'role', 
        'contact_number', 'address', 'discipline'
    ]
    
    # Style headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Add headers with styling
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Add sample data
    sample_data = [
        ['John', 'Doe', 'Robert Doe', 'john.doe@example.com', 'HOD', '1234567890', '123 Main St, City', 'Computer Science'],
        ['Jane', 'Smith', 'Michael Smith', 'jane.smith@example.com', 'Coordinator', '0987654321', '456 Oak Ave, Town', 'Mathematics'],
        ['Bob', 'Johnson', 'William Johnson', 'bob.johnson@example.com', 'Section Head', '5555555555', '789 Pine Rd, Village', 'Physics'],
    ]
    
    # Add sample data
    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = min(adjusted_width, 30)
    
    # Add instructions at the bottom
    instruction_row = len(sample_data) + 3
    ws.cell(row=instruction_row, column=1, value="INSTRUCTIONS:").font = Font(bold=True)
    ws.cell(row=instruction_row + 1, column=1, value="1. Role must be one of: HOD, Coordinator, Section Head")
    ws.cell(row=instruction_row + 2, column=1, value="2. Discipline names must match existing disciplines in system")
    ws.cell(row=instruction_row + 3, column=1, value="3. Contact Number and Address are optional")
    ws.cell(row=instruction_row + 4, column=1, value="4. Email must be unique")
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="admin_profile_sample.xlsx"'
    
    wb.save(response)
    return response
    from django.http import JsonResponse
from django.contrib.auth.decorators import login_required

@login_required
def debug_admin_permissions(request):
    """Debug view to check admin permissions"""
    user = request.user
    
    # Check if user has admin profile
    try:
        admin_profile = AdminProfile.objects.get(email=user.email)
        has_admin_profile = True
        admin_role = admin_profile.role
    except AdminProfile.DoesNotExist:
        has_admin_profile = False
        admin_role = None
    
    # Get all user attributes
    user_attrs = {
        'username': user.username,
        'email': user.email,
        'is_superuser': user.is_superuser,
        'is_staff': user.is_staff,
        'is_active': user.is_active,
        'is_admin': getattr(user, 'is_admin', False),
        'has_admin_profile': has_admin_profile,
        'admin_role': admin_role,
    }
    
    # Get all available attributes
    all_attrs = {}
    for attr in dir(user):
        if not attr.startswith('_') and not callable(getattr(user, attr)):
            try:
                all_attrs[attr] = str(getattr(user, attr))
            except:
                all_attrs[attr] = "Error getting value"
    
    return JsonResponse({
        'user_attributes': user_attrs,
        'all_attributes': all_attrs,
    })